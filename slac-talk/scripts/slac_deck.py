"""SlacDeck: build talks on the SLAC/Stanford template in James Baxter's style.

Import it from a per-talk build script (keep that script next to the talk, so the deck can be rebuilt):

    import sys, os
    sys.path.insert(0, os.path.expanduser("~/.claude/skills/slac-talk/scripts"))
    from slac_deck import SlacDeck, CARDINAL, BLACK

    d = SlacDeck()                                  # bundled FY27 template, its sample slides removed
    d.cover("Title", "Proposal category / meeting", "James Baxter", "Collaborators · date")
    d.build("Background: the partiality problem", [
        lambda s: d.picture(s, "figs/still.png", 0.88, 1.2, 5.6, 5.0),
        lambda s: d.text(s, 6.8, 1.3, 5.6, 3.0, ["-First point", "-Second point"]),
        lambda s: d.takeaway(s, "But it needs a lot of sample"),
    ], notes=["...", "...", "..."], source="Journal 10, 116 (2020)")
    d.save("talk.pptx")

Geometry is in inches on a 13.33 × 7.5 in slide. Content area: x 0.88–12.42, y 1.2–6.7; the bottom strip
(y ≥ 6.85) holds the SLAC logo, the citation line (x 1.79) and the slide number.
"""
import copy
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "SLAC_template_FY27.pptx")

CARDINAL = RGBColor(0x8C, 0x15, 0x15)   # Stanford cardinal: titles, table headers, stat numbers
ACCENT = RGBColor(0xB8, 0x3A, 0x4B)     # template accent: bullets, title rule, callout outline
TAKEAWAY = RGBColor(0xC0, 0x00, 0x00)   # red takeaway line at the end of a build
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x59, 0x59, 0x59)       # captions, citations
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_RED = RGBColor(0xD9, 0x7A, 0x86)   # secondary series in charts
LIGHT_GREY = RGBColor(0xA6, 0xA6, 0xA6)

CONTENT_LAYOUT = "Text Slide - 1 Column"
COVER_LAYOUT = "1_Title Slide - Option 4C - Printable"


class SlacDeck:
    def __init__(self, template=TEMPLATE):
        self.prs = Presentation(template)
        for sid in list(self.prs.slides._sldIdLst):  # drop the template's guidance slides
            self.prs.part.drop_rel(sid.rId)
            self.prs.slides._sldIdLst.remove(sid)
        layouts = {l.name: l for m in self.prs.slide_masters for l in m.slide_layouts}
        self.content_layout = layouts[CONTENT_LAYOUT]
        self.cover_layout = layouts[COVER_LAYOUT]
        self._sldnum = next(p for p in self.content_layout.placeholders if p.placeholder_format.idx == 4)
        self.last_build = []

    # ── slides ───────────────────────────────────────────────────────────────
    @staticmethod
    def ph(slide, idx):
        return next((p for p in slide.placeholders if p.placeholder_format.idx == idx), None)

    @staticmethod
    def remove(shape):
        shape._element.getparent().remove(shape._element)

    def cover(self, title, category, presenter, details, notes=None):
        """Grey 'Option 4C - Printable' cover. Reuse it as a section divider for long talks."""
        s = self.prs.slides.add_slide(self.cover_layout)
        for idx, val in ((10, title), (11, category), (12, presenter), (13, details)):
            self.ph(s, idx).text = val
        # The layout's headline box is bottom-anchored and short, so a 2–3 line title runs into the category line.
        # Give it explicit geometry (setting only .top on an inherited placeholder zeroes the other three) and top-anchor it.
        head = self.ph(s, 10)
        head.left, head.top, head.width, head.height = Inches(0.88), Inches(0.55), Inches(10.25), Inches(2.45)
        head.text_frame._txBody.bodyPr.set("anchor", "t")
        if len(title) > 40:
            for p in head.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(44)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def content(self, title, subtitle=None):
        """'Text Slide - 1 Column': cardinal title with the template rule; body placeholder removed."""
        s = self.prs.slides.add_slide(self.content_layout)
        s.shapes.title.text = title
        sub = self.ph(s, 13)
        if subtitle:
            sub.text = subtitle
        else:
            self.remove(sub)
        self.remove(self.ph(s, 16))
        s.shapes._spTree.append(copy.deepcopy(self._sldnum._element))  # python-pptx skips the slide number
        return s

    def build(self, title, steps, notes=None, subtitle=None, source=None):
        """Progressive build: one slide per step, slide k shows steps[0..k]. Returns the slides."""
        self.last_build = []
        for k in range(len(steps)):
            s = self.content(title, subtitle)
            for step in steps[: k + 1]:
                step(s)
            if source:
                self.source(s, source)
            if notes:
                s.notes_slide.notes_text_frame.text = notes[min(k, len(notes) - 1)]
            self.last_build.append(s)
        return self.last_build

    def pitch(self, quadrants, notes=None):
        """Elevator pitch on the cover background, a cardinal cross dividing four quadrants revealed one at a time.

        quadrants: four (header, bold_lead, rest, image_path_or_None) tuples, in the order
        question/problem → approach → result/impact → next/outlook. Image slots are small (top: 6.2 × 1.75 in,
        bottom: ~3 × 2.8 in): use single-panel or illustrative images, cropping multi-panel plots with `crop`.
        """
        boxes = [  # (text x, y, w, h), (image x, y, w, h)
            ((0.3, 0.2, 6.15, 1.45), (0.3, 1.7, 6.15, 1.75)),
            ((6.8, 0.2, 6.3, 1.45), (6.8, 1.7, 6.3, 1.75)),
            ((0.3, 3.65, 3.2, 2.9), (3.55, 3.7, 3.0, 2.8)),
            ((6.8, 3.65, 3.5, 2.9), (10.35, 3.7, 2.75, 2.8)),
        ]
        slides = []
        for k in range(len(quadrants)):
            s = self.prs.slides.add_slide(self.cover_layout)
            for p in list(s.placeholders):
                self.remove(p)
            for x1, y1, x2, y2 in ((6.62, 0.2, 6.62, 6.55), (0.2, 3.55, 13.1, 3.55)):
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
                c.line.color.rgb = CARDINAL
                c.line.width = Pt(1.5)
            for (head, lead, rest, img), (tbox, ibox) in list(zip(quadrants, boxes))[: k + 1]:
                self.text(s, *tbox, [[(head, True)], [(lead + " ", True), (rest, False)]], size=14)
                if img:
                    self.picture(s, img, *ibox)
            if notes:
                s.notes_slide.notes_text_frame.text = notes[min(k, len(notes) - 1)]
            slides.append(s)
        return slides

    def process(self, slide, steps, x=0.88, y=1.2, w=11.54, h=1.4, band=RGBColor(0xF6, 0xE7, 0xE9)):
        """Phase / method band: a tinted strip of rounded boxes joined by arrows (FY27 "Project overview").
        steps: [(title, detail), ...]; 3–5 fit across the slide. Returns the box shapes."""
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        bg.fill.solid()
        bg.fill.fore_color.rgb = band
        bg.line.fill.background()
        bg.shadow.inherit = False
        n, gap, pad = len(steps), 0.45, 0.15
        bw = (w - 2 * pad - gap * (n - 1)) / n
        boxes = []
        for i, (head, detail) in enumerate(steps):
            bx = x + pad + i * (bw + gap)
            b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(y + pad), Inches(bw), Inches(h - 2 * pad))
            b.fill.solid()
            b.fill.fore_color.rgb = WHITE
            b.line.color.rgb = ACCENT
            b.line.width = Pt(1)
            b.shadow.inherit = False
            tf = b.text_frame
            tf.word_wrap = True
            for j, (t, bold, size, color) in enumerate(((head, True, 13, CARDINAL), (detail, False, 11, BLACK))):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = t
                r.font.bold, r.font.size, r.font.color.rgb = bold, Pt(size), color
            boxes.append(b)
            if i:
                ax = bx - gap
                c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ax + 0.06), Inches(y + h / 2),
                                               Inches(bx - 0.06), Inches(y + h / 2))
                c.line.color.rgb = CARDINAL
                c.line.width = Pt(1.5)
                etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"), type="triangle")
        return boxes

    @staticmethod
    def crop(path, out, box):
        """Crop a figure before placing it, e.g. one panel of a multi-panel notebook plot (wide strips come out
        unreadably small in pitch quadrants and half-width slots). box = (left, top, right, bottom) as fractions."""
        im = Image.open(path)
        W, H = im.size
        l, t, r, b = box
        im.crop((int(W * l), int(H * t), int(W * r), int(H * b))).save(out)
        return out

    def figure_slide(self, title, image, caption=None, source=None, notes=None):
        """One large figure with a caption: the backup-slide pattern."""
        s = self.content(title)
        self.picture(s, image, 0.88, 1.2, 11.54, 5.0)
        if caption:
            self.caption(s, caption, 0.88, 6.25, 11.54, size=13)
        if source:
            self.source(s, source)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # ── elements ─────────────────────────────────────────────────────────────
    @staticmethod
    def picture(slide, path, x, y, w, h, align="center"):
        """Fit an image inside the box, keeping its aspect ratio."""
        iw, ih = Image.open(path).size
        scale = min(w / iw, h / ih)
        pw, phh = iw * scale, ih * scale
        px = x + (w - pw) / 2 if align == "center" else x
        py = y + (h - phh) / 2 if align == "center" else y
        return slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(phh))

    @staticmethod
    def _bullet(p):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "228600")
        pPr.set("indent", "-228600")
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr"), val="B83A4B")
        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(pPr, qn("a:buChar"), char="•")

    def text(self, slide, x, y, w, h, paras, size=16, color=BLACK, align=PP_ALIGN.LEFT, bullets=False):
        """paras: str (a leading '-' makes it a bullet) or a list of (text, bold) runs."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            bullet = False
            if isinstance(para, str):
                if para.startswith("-"):
                    bullet, para = True, para[1:].strip()
                runs = [(para, False)]
            else:
                runs = para
            for t, b in runs:
                r = p.add_run()
                r.text = t
                r.font.size = Pt(size)
                r.font.bold = b
                r.font.color.rgb = color
            p.space_after = Pt(4)
            if bullet or bullets:
                self._bullet(p)
        return tb

    def labelled_list(self, slide, x, y, w, h, sections, size=16):
        """Conclusions style: cardinal bold section labels (18 pt) with bullets under each.
        sections: [(label, [bullet, ...]), ...]."""
        paras = []
        for label, items in sections:
            paras.append([(label, True)])
            paras += ["-" + it for it in items]
        tb = self.text(slide, x, y, w, h, paras, size=size)
        for p in tb.text_frame.paragraphs:
            if p.runs and p.runs[0].font.bold:
                for r in p.runs:
                    r.font.color.rgb = CARDINAL
                    r.font.size = Pt(size + 2)
        return tb

    def takeaway(self, slide, msg, y=6.2, x=0.88, w=11.54, h=0.55, size=18):
        """The red line that ends a build."""
        return self.text(slide, x, y, w, h, [[(msg, False)]], size=size, color=TAKEAWAY)

    def callout(self, slide, msg, x, y, w, h, size=16, bold_lead=None):
        """White box with an accent outline; bold cardinal lead ('But:', 'Open question:', 'Next test:')."""
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = ACCENT
        box.line.width = Pt(1.5)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if bold_lead:
            r = p.add_run()
            r.text = bold_lead + " "
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = CARDINAL
        r = p.add_run()
        r.text = msg
        r.font.size = Pt(size)
        r.font.color.rgb = BLACK
        return box

    def source(self, slide, msg):
        """Small citation next to the SLAC logo ('Journal vol, page (year)' or 'notebook · NOTES §')."""
        return self.text(slide, 1.79, 6.87, 7.4, 0.4, [msg], size=11, color=GREY)

    def caption(self, slide, msg, x, y, w, size=12):
        return self.text(slide, x, y, w, 0.4, [msg], size=size, color=GREY, align=PP_ALIGN.CENTER)

    def stat(self, slide, big, label, x, y=1.3, w=3.6):
        """Stat card: big cardinal number, label underneath. Three fit across at x = 0.9, 4.85, 8.8."""
        size = 40 if len(big) <= 8 else 28  # long values ("0.080° → 0.056°") must stay on one line
        self.text(slide, x, y, w, 1.0, [[(big, True)]], size=size, color=CARDINAL, align=PP_ALIGN.CENTER)
        self.text(slide, x, y + 1.05, w, 0.9, [label], size=14, align=PP_ALIGN.CENTER)

    def table(self, slide, x, y, w, rows, header, widths=None, highlight=None, size=14, left_cols=(0,)):
        """Native table: cardinal header, plain white body; `highlight` = body row index in bold cardinal."""
        shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w),
                                       Inches(0.42 * (len(rows) + 1)))
        tbl = shape.table
        if widths:
            for j, cw in enumerate(widths):
                tbl.columns[j].width = Inches(cw)
        for i, row in enumerate([header] + list(rows)):
            hi = highlight is not None and i - 1 == highlight
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARDINAL if i == 0 else WHITE
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.bold = i == 0 or hi
                        r.font.color.rgb = WHITE if i == 0 else CARDINAL if hi else BLACK
        return shape

    @staticmethod
    def _style_chart(chart, title):
        chart.font.size = Pt(12)
        chart.font.name = "Lato"
        chart.has_title = True
        chart.chart_title.text_frame.text = title
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(13)
        run.font.bold = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    @staticmethod
    def _style_series(chart, series):
        for ps, (_, _, color, dashed) in zip(chart.plots[0].series, series):
            ps.smooth = False
            ps.format.line.color.rgb = color
            ps.format.line.width = Pt(1.5 if dashed else 2.25)
            if dashed:
                ps.format.line.dash_style = MSO_LINE.DASH
            ps.marker.style = XL_MARKER_STYLE.CIRCLE
            ps.marker.size = 7
            ps.marker.format.fill.solid()
            ps.marker.format.fill.fore_color.rgb = color
            ps.marker.format.line.color.rgb = color

    @staticmethod
    def _axis(ax, lo, hi, unit, fmt):
        if lo is not None:
            ax.minimum_scale = lo
        if hi is not None:
            ax.maximum_scale = hi
        if unit is not None:
            ax.major_unit = unit
        ax.tick_labels.number_format = fmt
        ax.tick_labels.number_format_is_linked = False

    def line_chart(self, slide, x, y, w, h, categories, series, title, ymin=None, ymax=None, yunit=None, yfmt="0.0"):
        """Native category line chart (e.g. CC per resolution bin). series = (name, values, color, dashed)."""
        data = CategoryChartData()
        data.categories = categories
        for name, vals, _, _ in series:
            data.add_series(name, vals)
        chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        self._style_chart(chart, title)
        self._axis(chart.value_axis, ymin, ymax, yunit, yfmt)
        chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        self._style_series(chart, series)
        return chart

    def scatter_chart(self, slide, x, y, w, h, series, title, xmin=None, xmax=None, xunit=None, xfmt="0",
                      ymin=None, ymax=None, yunit=None, yfmt="0.0"):
        """Native XY chart with lines; series = (name, ((x, y), ...), color, dashed). Use for numeric x
        (box size, divergence, N crystals) and for reference lines (a two-point dashed series)."""
        data = XyChartData()
        for name, pts, _, _ in series:
            sr = data.add_series(name)
            for px, py in pts:
                sr.add_data_point(px, py)
        chart = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        self._style_chart(chart, title)
        self._axis(chart.category_axis, xmin, xmax, xunit, xfmt)
        self._axis(chart.value_axis, ymin, ymax, yunit, yfmt)
        chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        self._style_series(chart, series)
        for ps, (_, _, _, dashed) in zip(chart.plots[0].series, series):
            if dashed:
                ps.marker.style = XL_MARKER_STYLE.NONE
        return chart

    def save(self, path):
        self.prs.save(path)
        print("wrote", path, len(self.prs.slides), "slides")
        return path
